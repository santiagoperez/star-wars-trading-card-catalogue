#!/usr/bin/env python3
"""Import card rows from an exported Google Sheets workbook."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "Missing dependency: install it with "
        "'python -m pip install -r requirements.txt'."
    )
    raise SystemExit(2)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "data" / "cards"
REQUIRED_COLUMNS = (
    "Year",
    "Collection",
    "#",
    "Title",
)
OPTIONAL_SCALAR_COLUMNS = {
    "Section": "section",
    "Subtitle": "subtitle",
    "Franchise": "franchise",
    "Insert Collection": "insert_collection",
    "Description": "description",
    "Front description": "front_description",
    "Front Description": "front_description",
    "Back description": "back_description",
    "Back Description": "back_description",
    "Species": "species",
    "Home World": "home_world",
    "Location": "location",
}


def clean_text(value: object) -> str:
    """Convert a spreadsheet cell to trimmed text without changing its wording."""
    return "" if value is None else str(value).strip()


def whole_number(value: object, field: str, row_number: int) -> int:
    """Read a spreadsheet numeric cell that must contain a whole number."""
    if isinstance(value, bool):
        raise ValueError(f"Row {row_number}: {field} must be a whole number")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Row {row_number}: invalid {field}: {value!r}") from error
    if not number.is_integer():
        raise ValueError(f"Row {row_number}: {field} must be a whole number")
    return int(number)


def card_number(value: object, row_number: int) -> str:
    """Preserve text card numbers while removing Excel's decimal from integers."""
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Row {row_number}: invalid card number: {value!r}")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = clean_text(value)
    if not result:
        raise ValueError(f"Row {row_number}: card number is empty")
    return result


def slug(value: str) -> str:
    """Create a stable lowercase ASCII path/id component."""
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", "-", ascii_value).strip("-")


def split_values(value: object) -> list[str]:
    """Split slash-separated source values while preserving their wording."""
    text = clean_text(value)
    if not text:
        return []
    return list(dict.fromkeys(part.strip() for part in text.split("/") if part.strip()))


def make_card(
    row: dict[str, object],
    row_number: int,
    allow_blank_title: bool = False,
) -> tuple[dict[str, object], str]:
    year = whole_number(row["Year"], "year", row_number)
    collection = clean_text(row["Collection"])
    number = card_number(row["#"], row_number)
    title = clean_text(row["Title"])
    if not collection:
        raise ValueError(f"Row {row_number}: collection is required")
    if not title and not allow_blank_title:
        raise ValueError(f"Row {row_number}: title is required")

    collection_slug = slug(collection)
    number_slug = slug(number)
    if not collection_slug or not number_slug:
        raise ValueError(f"Row {row_number}: collection or card number cannot form an id")

    padded_number = number.zfill(3) if number.isdigit() else number_slug
    card: dict[str, object] = {
        "id": f"{year}-{collection_slug}-{padded_number.lower()}",
        "year": year,
        "collection": collection,
        "number": number,
        "title": title,
    }

    for source_name, field_name in OPTIONAL_SCALAR_COLUMNS.items():
        value = clean_text(row.get(source_name))
        if value and field_name not in card:
            card[field_name] = value

    affiliation = split_values(row.get("Affiliation"))
    if affiliation:
        card["affiliation"] = affiliation

    notes = clean_text(row.get("Notes"))
    if notes:
        card["notes"] = [notes]

    return card, f"{padded_number}.json"


def import_sheet(
    workbook_path: Path,
    sheet_name: str,
    output_root: Path,
    overwrite: bool = False,
    dry_run: bool = False,
    template_incomplete: bool = False,
) -> int:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {available}")

    rows = workbook[sheet_name].iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as error:
        raise ValueError(f"Sheet {sheet_name!r} is empty") from error

    headers = [clean_text(value) for value in raw_headers]
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"Sheet {sheet_name!r} is missing columns: {', '.join(missing)}")
    planned: list[tuple[Path, dict[str, object]]] = []
    ids: set[str] = set()
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and clean_text(value) for value in values):
            continue
        source = {
            name: values[index] if index < len(values) else None
            for index, name in enumerate(headers)
            if name
        }
        incomplete = not clean_text(source.get("Title"))
        card, filename = make_card(
            source,
            row_number,
            allow_blank_title=template_incomplete and incomplete,
        )
        if card["id"] in ids:
            raise ValueError(f"Row {row_number}: duplicate card id {card['id']!r}")
        ids.add(str(card["id"]))
        if incomplete:
            filename += ".template"
        destination = output_root / str(card["year"]) / slug(str(card["collection"])) / filename
        planned.append((destination, card))

    collisions = [path for path, _ in planned if path.exists() and not overwrite]
    if collisions:
        preview = ", ".join(str(path) for path in collisions[:3])
        suffix = " ..." if len(collisions) > 3 else ""
        raise FileExistsError(f"Refusing to overwrite {len(collisions)} file(s): {preview}{suffix}")

    if not dry_run:
        for destination, card in planned:
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_text(
                json.dumps(card, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )

    template_count = sum(path.name.endswith(".json.template") for path, _ in planned)
    card_count = len(planned) - template_count
    action = "Would import" if dry_run else "Imported"
    template_action = "would create" if dry_run else "created"
    print(
        f"{action} {card_count} card(s) and {template_action} "
        f"{template_count} template(s) from {sheet_name!r}."
    )
    return len(planned)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Exported .xlsx workbook path.")
    parser.add_argument("--sheet", required=True, help="Exact worksheet tab name.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Card data root.")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing card files.")
    parser.add_argument("--dry-run", action="store_true", help="Check input without writing files.")
    parser.add_argument(
        "--template-incomplete",
        action="store_true",
        help="Write rows with blank titles as .json.template files.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        import_sheet(
            args.workbook,
            args.sheet,
            args.output,
            args.overwrite,
            args.dry_run,
            args.template_incomplete,
        )
    except (FileNotFoundError, OSError, ValueError) as error:
        print(f"Import failed: {error}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
